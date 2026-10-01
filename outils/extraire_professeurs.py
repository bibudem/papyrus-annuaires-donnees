#!/usr/bin/env python3
"""
Extrait les professeurs d'un fichier .txt à LARGEUR FIXE (export Synchro),
affiche un tableau Nom / Prénom / Courriel / CodeUnite / Statut, et ajoute
l'ORCID (correspondance par courriel) si un fichier Excel est fourni.

Usage:
    python extraire_professeurs.py fichier.txt
    python extraire_professeurs.py fichier.txt --diagnostic
    python extraire_professeurs.py fichier.txt --excel annuaire.xlsx
    python extraire_professeurs.py fichier.txt --excel annuaire.xlsx --csv sortie.csv

Dépendance pour --excel : openpyxl (pip install openpyxl si absent).
"""

import argparse
import csv
import re
import sys

# =============================================================================
# CONFIGURATION
# Tout ce qui est propre au format du fichier Synchro ou à l'affichage est
# regroupé ici. C'est la seule section à modifier si le format du fichier
# change ou si de nouveaux titres/colonnes doivent être reconnus.
# =============================================================================

# --- Lecture du fichier ---
# Encodages essayés dans l'ordre (les exports Windows sont souvent en
# cp1252/latin-1 plutôt qu'en UTF-8).
ENCODAGES_A_ESSAYER = ["utf-8", "cp1252", "latin-1"]

# --- Position des champs dans la ligne à largeur fixe ---
# (debut, fin) -- indices de caractères, 0-indexé, fin exclue.
# Nom/Prénom est à position fixe. Courriel et CodeUnite sont localisés
# dynamiquement dans ZONE_COURRIEL (voir extraire_courriel_et_code_unite)
# car leur position exacte peut varier légèrement d'une ligne à l'autre.
CHAMP_NOM_PRENOM = (0, 40)
ZONE_COURRIEL = (166, 320)

# --- Titres considérés comme "Professeur" ---
# Tirés du filtre Synchro sur la colonne "Fonction". Inclut les formes
# abrégées ("Prof ...") qui ne contiennent pas le mot complet "professeur".
# Pour ajouter/retirer un titre : modifier cette liste seulement: "Professeur invité",
TITRES_PROFESSEUR = [
    "Prof agrégé(e) PTG SC Acad.",
    "Prof agrégé(e) PTG SC Recherc.",
    "Prof titulaire PTG SC Acad.",
    "Prof. sous octroi titulaire",
    "Professeur associé",
    "Professeur form. prat. titulai",
    "Professeur sous octroi adjoint",
    "Professeur(e) adjoint(e)",
    "Professeur(e) agrégé(e)",
    "Professeur(e) form.prat agrég.",
    "Professeur(e) sous oct. agrég.",
    "Professeur(e) titulaire",
    "Professeur/ chercheur adjoint",
    "Professeur/ chercheur agrégé",
    "Professeur/chercheur titulaire",
]

# --- Colonnes recherchées dans le fichier Excel (par mot-clé d'en-tête) ---
MOT_CLE_COLONNE_COURRIEL = "courriel"
MOT_CLE_COLONNE_ORCID = "orcid"

# --- Colonnes du tableau de résultat, dans l'ordre d'affichage ---
ENTETES_SORTIE = ["Nom", "Prénom", "Courriel", "CodeUnite", "Statut"]
ENTETE_ORCID = "ORCID"  # ajoutée seulement si --excel est fourni

# --- Export CSV ---
CSV_SEPARATEUR = ";"       # ';' pour compatibilité Excel FR
CSV_ENCODAGE = "utf-8-sig"  # BOM pour qu'Excel détecte l'UTF-8 correctement

# --- Diagnostic (--diagnostic) ---
DIAGNOSTIC_NB_LIGNES_EXEMPLE = 3
DIAGNOSTIC_REGLE_DEBUT = 140  # bornes de la "règle de positions" affichée
DIAGNOSTIC_REGLE_FIN = 310    # pour repérer un champ à la main

# =============================================================================
# MOTIFS (regex précompilées à partir de la configuration ci-dessus)
# =============================================================================


def _construire_motif_titre(titre):
    """Regex insensible à la casse, tolérant un nombre variable d'espaces
    entre les mots du titre (padding de largeur fixe)."""
    return re.compile(r"\s+".join(re.escape(mot) for mot in titre.split()), re.IGNORECASE)


MOTIFS_PROFESSEUR = [_construire_motif_titre(titre) for titre in TITRES_PROFESSEUR]
MOTIF_COURRIEL = re.compile(r"\S+@\S+")
MOTIF_ORCID = re.compile(r"\d{4}-\d{4}-\d{4}-\d{3}[\dXx]")


# =============================================================================
# LECTURE ET EXTRACTION DU FICHIER TXT
# =============================================================================

def lire_lignes(chemin_fichier):
    """Lit le fichier ligne par ligne en essayant chaque encodage de
    ENCODAGES_A_ESSAYER jusqu'à ce que l'un fonctionne."""
    for encodage in ENCODAGES_A_ESSAYER:
        try:
            with open(chemin_fichier, encoding=encodage) as f:
                lignes = [l.rstrip("\r\n") for l in f]
            return [l for l in lignes if l.strip()]
        except UnicodeDecodeError:
            continue
    raise UnicodeDecodeError("Impossible de décoder le fichier avec les encodages testés.")


def extraire_courriel_et_code_unite(ligne):
    """Localise le courriel par son contenu (motif '@') dans ZONE_COURRIEL,
    puis prend le bloc de texte juste avant comme CodeUnite. Plus robuste
    qu'une position fixe, car ça s'ajuste même si l'alignement varie."""
    debut_zone, fin_zone = ZONE_COURRIEL
    zone = ligne[debut_zone:fin_zone]

    m = MOTIF_COURRIEL.search(zone)
    if not m:
        return "", ""

    courriel = m.group(0).strip()
    if courriel and courriel[0].isupper():
        courriel = courriel[1:]  # lettre parasite collée devant le vrai courriel

    debut_absolu = debut_zone + m.start()
    avant = ligne[max(0, debut_absolu - 40):debut_absolu]
    tokens = re.findall(r"\S+", avant)
    code_unite = tokens[-1] if tokens else ""

    return courriel, code_unite


def extraire_infos(ligne):
    """Extrait Nom, Prénom, Courriel, CodeUnite d'une ligne (sans le Statut,
    qui vient du titre trouvé par filtrer_professeurs)."""
    debut, fin = CHAMP_NOM_PRENOM
    nom_prenom = ligne[debut:fin].strip()
    nom, _, prenom = nom_prenom.partition(",")

    courriel, code_unite = extraire_courriel_et_code_unite(ligne)
    return [nom.strip(), prenom.strip(), courriel, code_unite]


def filtrer_professeurs(lignes):
    """Retourne les lignes correspondant à un titre de TITRES_PROFESSEUR,
    sous forme de tuples (ligne, statut) où statut est le texte EXACT tel
    qu'écrit dans le fichier (casse et espacement d'origine)."""
    resultat = []
    for ligne in lignes:
        for motif in MOTIFS_PROFESSEUR:
            m = motif.search(ligne)
            if m:
                resultat.append((ligne, m.group(0)))
                break
    return resultat


# =============================================================================
# CORRESPONDANCE ORCID VIA FICHIER EXCEL
# =============================================================================

def normaliser_courriel(courriel):
    return courriel.strip().lower() if courriel else ""


def trouver_colonne(entetes, mot_cle, eviter_index=None):
    """Trouve l'index de la première colonne dont l'en-tête contient mot_cle
    (insensible à la casse), en évitant optionnellement une colonne déjà prise."""
    for i, valeur in enumerate(entetes):
        if i != eviter_index and valeur and mot_cle in str(valeur).strip().lower():
            return i
    return None


def extraire_premier_orcid(valeur_cellule):
    """Extrait le premier ORCID valide (format 0000-0000-0000-0000) d'une
    cellule, même si plusieurs y sont collés (ex: '0000-... / 0000-...').
    Retourne "" si la cellule ne contient aucun ORCID valide (ex: 'N/A',
    'à venir') : mieux vaut pas d'ORCID qu'un texte invalide dans Papyrus."""
    if not valeur_cellule:
        return ""
    m = MOTIF_ORCID.search(str(valeur_cellule))
    return m.group(0).upper() if m else ""


def charger_orcid_depuis_excel(chemin_excel):
    """Lit le fichier Excel et retourne {courriel_normalise: [orcid1, orcid2, ...]}.

    Les colonnes Courriel et ORCID sont trouvées par leur en-tête (via
    MOT_CLE_COLONNE_COURRIEL / MOT_CLE_COLONNE_ORCID), pas par lettre fixe.
    Si un courriel a plusieurs ORCID distincts, les deux sont conservés (et
    signalés) plutôt que d'écraser l'un des deux en silence — voir main()
    pour la règle qui choisit lequel afficher.

    Lève RuntimeError (plutôt que de quitter) si openpyxl manque ou si les
    colonnes sont introuvables, pour que l'appelant puisse journaliser l'erreur.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("le module 'openpyxl' est requis pour lire le fichier Excel "
                           "(pip install openpyxl).")

    lignes = list(load_workbook(chemin_excel, data_only=True).active.iter_rows(values_only=True))
    if not lignes:
        return {}

    entetes = lignes[0]
    idx_courriel = trouver_colonne(entetes, MOT_CLE_COLONNE_COURRIEL)
    idx_orcid = trouver_colonne(entetes, MOT_CLE_COLONNE_ORCID, eviter_index=idx_courriel)

    if idx_courriel is None or idx_orcid is None:
        raise RuntimeError(f"impossible de trouver les colonnes 'Courriel' et/ou 'ORCID' dans "
                           f"'{chemin_excel}'. En-têtes trouvées : {list(entetes)}")

    correspondance = {}
    nb_orcid_invalides = 0
    for ligne in lignes[1:]:
        if idx_courriel >= len(ligne) or idx_orcid >= len(ligne):
            continue
        courriel = normaliser_courriel(ligne[idx_courriel])
        orcid = extraire_premier_orcid(ligne[idx_orcid])
        if not orcid and str(ligne[idx_orcid] or "").strip():
            nb_orcid_invalides += 1
        if courriel and orcid:
            liste = correspondance.setdefault(courriel, [])
            if orcid not in liste:
                liste.append(orcid)

    if nb_orcid_invalides:
        print(f"⚠ ATTENTION : {nb_orcid_invalides} cellule(s) ORCID non vide(s) mais sans ORCID "
              f"valide (format 0000-0000-0000-000X) — ignorée(s).")

    doublons = {c: o for c, o in correspondance.items() if len(o) > 1}
    if doublons:
        print(f"⚠ ATTENTION : {len(doublons)} courriel(s) avec plusieurs ORCID différents "
              f"(seul le premier est utilisé) :")
        for courriel, orcids in doublons.items():
            print(f"    {courriel} -> {' / '.join(orcids)}")
        print()

    return correspondance


# =============================================================================
# AFFICHAGE (tableau, export CSV, diagnostic)
# =============================================================================

def ecrire_csv(chemin_sortie, entetes, lignes_extraites):
    """Écrit le résultat dans un fichier CSV (voir CSV_SEPARATEUR / CSV_ENCODAGE)."""
    with open(chemin_sortie, "w", newline="", encoding=CSV_ENCODAGE) as f:
        writer = csv.writer(f, delimiter=CSV_SEPARATEUR)
        writer.writerow(entetes)
        writer.writerows(lignes_extraites)


def afficher_tableau(entetes, lignes_extraites):
    """Affiche un tableau aligné dans la console."""
    if not lignes_extraites:
        print("Aucun professeur trouvé.")
        return

    largeurs = [
        max(len(entetes[i]), max((len(l[i]) for l in lignes_extraites), default=0))
        for i in range(len(entetes))
    ]

    def formater_ligne(champs):
        return " | ".join(champ.ljust(largeurs[i]) for i, champ in enumerate(champs))

    print(formater_ligne(entetes))
    print("-+-".join("-" * largeur for largeur in largeurs))
    for ligne in lignes_extraites:
        print(formater_ligne(ligne))


def afficher_regle(ligne, debut=DIAGNOSTIC_REGLE_DEBUT, fin=DIAGNOSTIC_REGLE_FIN):
    """Affiche les numéros de colonnes au-dessus d'une tranche de la ligne,
    pour repérer visuellement une position sans partager de vraies données."""
    tranche = ligne[debut:fin]
    dizaines = "".join(str((debut + i) // 10 % 10) for i in range(len(tranche)))
    unites = "".join(str((debut + i) % 10) for i in range(len(tranche)))
    print(f"  Position de départ de la tranche : {debut}")
    print(f"  {dizaines}")
    print(f"  {unites}")
    print(f"  {tranche}")


def afficher_diagnostic(lignes, nb_lignes_exemple=DIAGNOSTIC_NB_LIGNES_EXEMPLE):
    """Affiche les champs extraits + une règle de positions, pour validation
    manuelle du format (utile quand le fichier source change de structure)."""
    print("=" * 70)
    print("DIAGNOSTIC (largeur fixe)")
    print("=" * 70)
    print(f"Longueur de ligne détectée : {len(lignes[0]) if lignes else 0} caractères\n")

    for n, ligne in enumerate(lignes[:nb_lignes_exemple]):
        nom, prenom, courriel, code_unite = extraire_infos(ligne)
        print(f"Ligne {n + 1} (longueur={len(ligne)}) :")
        print(f"  Nom        : {nom!r}")
        print(f"  Prénom     : {prenom!r}")
        print(f"  Courriel   : {courriel!r}")
        print(f"  CodeUnite  : {code_unite!r}")
        print()
        print("  Règle de positions (pour repérer le champ CodeUnite à la main) :")
        afficher_regle(ligne)
        print()

    print("=" * 70 + "\n")


# =============================================================================
# POINT D'ENTRÉE
# =============================================================================

def main():
    parser = argparse.ArgumentParser(
        description="Extrait les professeurs d'un fichier txt à largeur fixe, "
                     "avec correspondance ORCID optionnelle via un fichier Excel."
    )
    parser.add_argument("fichier_txt", help="Chemin vers le fichier .txt source")
    parser.add_argument("--diagnostic", action="store_true",
                         help="Affiche la structure du fichier au lieu du tableau")
    parser.add_argument("--excel", metavar="FICHIER.xlsx",
                         help="Fichier Excel contenant les ORCID (colonnes 'Courriel' et 'ORCID')")
    parser.add_argument("--csv", metavar="SORTIE.csv",
                         help="Écrit aussi le résultat dans un fichier CSV")
    args = parser.parse_args()

    try:
        lignes = lire_lignes(args.fichier_txt)
    except FileNotFoundError:
        print(f"Erreur : le fichier '{args.fichier_txt}' est introuvable.")
        sys.exit(1)

    if args.diagnostic:
        afficher_diagnostic(lignes)
        return

    professeurs = filtrer_professeurs(lignes)
    print(f"Total de professeurs trouvés : {len(professeurs)}\n")

    lignes_extraites = [extraire_infos(ligne) + [statut] for ligne, statut in professeurs]
    entetes = list(ENTETES_SORTIE)

    if args.excel:
        try:
            orcid_par_courriel = charger_orcid_depuis_excel(args.excel)
        except FileNotFoundError:
            print(f"Erreur : le fichier Excel '{args.excel}' est introuvable.")
            sys.exit(1)
        except RuntimeError as e:
            print(f"Erreur : {e}")
            sys.exit(1)

        entetes.append(ENTETE_ORCID)
        trouves = 0
        for infos in lignes_extraites:
            orcids = orcid_par_courriel.get(normaliser_courriel(infos[2]), [])
            trouves += bool(orcids)
            infos.append(orcids[0] if orcids else "")  # seul le 1er ORCID est gardé

        print(f"ORCID trouvés : {trouves} / {len(lignes_extraites)}\n")

    if args.csv:
        ecrire_csv(args.csv, entetes, lignes_extraites)
        print(f"Fichier CSV écrit : {args.csv}\n")

    afficher_tableau(entetes, lignes_extraites)


if __name__ == "__main__":
    main()
